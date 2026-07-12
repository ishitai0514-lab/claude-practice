#!/usr/bin/env python3
"""マッチング判定結果を results.xlsx に出力する。

候補者ごとに1シート（氏名をシート名にする）を作成し、
求人名・A軸・A軸理由・B軸・B軸理由・総合判断の列を持つ。
総合判断が「本命として提案」の行は色付けする。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from match import ASSESSMENTS, CANDIDATE_IDS, JOB_IDS, load_candidates, load_jobs, ng_check, overall_judgment

REPO_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = REPO_ROOT / "results.xlsx"

HEADERS = ["求人名", "A軸", "A軸理由", "B軸", "B軸理由", "総合判断"]
COLUMN_WIDTHS = [40, 6, 60, 6, 60, 34]

HEADER_FILL = PatternFill(start_color="FFDCE6F1", end_color="FFDCE6F1", fill_type="solid")
HONMEI_FILL = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_TOP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")


def build_sheet(wb: Workbook, candidate_id: str, jobs: dict, cand: dict) -> None:
    sheet_name = cand.get("氏名", candidate_id)
    ws = wb.create_sheet(title=sheet_name)

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_TOP_CENTER

    row_idx = 2
    for jid in JOB_IDS:
        job = jobs[jid]
        label = f"{job.get('会社名', '')} / {job.get('ポジション名', '')}"
        ok, ng_reason = ng_check(candidate_id, jid)

        if not ok:
            values = [label, "-", f"対象外：{ng_reason}", "-", "-", "対象外"]
        else:
            key = (candidate_id, jid)
            assess = ASSESSMENTS.get(key)
            if assess is None:
                values = [label, "未評価", "-", "未評価", "-", "-"]
            else:
                judgment = overall_judgment(assess["a"], assess["b"])
                values = [label, assess["a"], assess["a_reason"], assess["b"], assess["b_reason"], judgment]

        is_honmei = values[-1] == "本命として提案"
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = WRAP_TOP_CENTER if col in (2, 4) else WRAP_TOP
            if is_honmei:
                cell.fill = HONMEI_FILL
        row_idx += 1

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"


def main() -> None:
    jobs = load_jobs()
    candidates = load_candidates()

    wb = Workbook()
    wb.remove(wb.active)

    for cid in CANDIDATE_IDS:
        build_sheet(wb, cid, jobs, candidates[cid])

    wb.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
